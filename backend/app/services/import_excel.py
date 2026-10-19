"""
Import historical data from user's Excel workbook
(Показания счётчиков / оплата ЖКХ format).
"""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any, Optional

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.meter import Meter, MeterReading
from app.models.bill import Bill, BillItem, Payment, ServiceType
from app.models.property import Property

MONTH_MAP = {
    "январ": 1, "феврал": 2, "март": 3, "апрел": 4, "ма": 5, "май": 5,
    "июн": 6, "июл": 7, "август": 8, "сентябр": 9, "октябр": 10,
    "ноябр": 11, "декабр": 12,
}


def _month_num(val: Any) -> Optional[int]:
    if val is None:
        return None
    if isinstance(val, (int, float)) and 1 <= int(val) <= 12:
        return int(val)
    s = str(val).strip().lower().replace("ё", "е")
    for prefix, num in MONTH_MAP.items():
        if s.startswith(prefix):
            return num
    return None


def _as_date(val: Any, year: int, month: int) -> date:
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, str):
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
            try:
                return datetime.strptime(val[:10], fmt).date()
            except ValueError:
                pass
    # mid-month fallback
    return date(year, month, min(15, 28))


def _num(val: Any) -> Optional[float]:
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).replace(" ", "").replace("\xa0", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _get_or_create_meter(
    db: Session,
    property_id: int,
    mtype: str,
    name: str,
    tariff_type: str = "single",
    unit: str = "м³",
) -> Meter:
    m = db.execute(
        select(Meter).where(
            Meter.property_id == property_id,
            Meter.type == mtype,
            Meter.name == name,
        )
    ).scalar_one_or_none()
    if m:
        return m
    m = Meter(
        property_id=property_id,
        type=mtype,
        name=name,
        tariff_type=tariff_type,
        unit=unit,
        is_active=True,
    )
    db.add(m)
    db.flush()
    return m


def _reading_exists(db: Session, meter_id: int, reading_date: date) -> bool:
    return (
        db.execute(
            select(MeterReading).where(
                MeterReading.meter_id == meter_id,
                MeterReading.reading_date == reading_date,
            )
        ).scalar_one_or_none()
        is not None
    )


def _bill_exists(db: Session, property_id: int, year: int, month: int, provider: str) -> Optional[Bill]:
    return db.execute(
        select(Bill).where(
            Bill.property_id == property_id,
            Bill.period_year == year,
            Bill.period_month == month,
            Bill.provider == provider,
        )
    ).scalar_one_or_none()


def _import_water_sheet(db: Session, ws, property_id: int, year_hint: int) -> dict:
    """Sheet like '2025': ХВ СУ1/СУ2 + ГВ СУ1/СУ2."""
    cold = _get_or_create_meter(db, property_id, "cold_water", "СУ1+СУ2 ХВ", "single", "м³")
    hot = _get_or_create_meter(db, property_id, "hot_water", "СУ1+СУ2 ГВ", "single", "м³")
    # Also individual counters if useful for display
    cold1 = _get_or_create_meter(db, property_id, "cold_water", "СУ1 ХВ", "single", "м³")
    cold2 = _get_or_create_meter(db, property_id, "cold_water", "СУ2 ХВ", "single", "м³")
    hot1 = _get_or_create_meter(db, property_id, "hot_water", "СУ1 ГВ", "single", "м³")
    hot2 = _get_or_create_meter(db, property_id, "hot_water", "СУ2 ГВ", "single", "м³")

    created = 0
    skipped = 0
    for row in ws.iter_rows(min_row=1, values_only=True):
        cells = list(row) if row else []
        if len(cells) < 7:
            continue
        # Find month in first few columns
        month = None
        col0 = None
        for i, c in enumerate(cells[:4]):
            month = _month_num(c)
            if month:
                col0 = i
                break
        if not month:
            continue
        # After month: Date, ХВ SU1, SU2, total, consumption, empty, ГВ SU1, SU2, total, consumption
        base = col0 + 1
        dt = cells[base] if base < len(cells) else None
        hv1 = _num(cells[base + 1] if base + 1 < len(cells) else None)
        hv2 = _num(cells[base + 2] if base + 2 < len(cells) else None)
        hv_tot = _num(cells[base + 3] if base + 3 < len(cells) else None)
        hv_cons = _num(cells[base + 4] if base + 4 < len(cells) else None)
        # skip empty col often at base+5
        gv1 = _num(cells[base + 6] if base + 6 < len(cells) else None)
        gv2 = _num(cells[base + 7] if base + 7 < len(cells) else None)
        gv_tot = _num(cells[base + 8] if base + 8 < len(cells) else None)
        gv_cons = _num(cells[base + 9] if base + 9 < len(cells) else None)

        if hv_tot is None and hv1 is not None:
            hv_tot = (hv1 or 0) + (hv2 or 0)
        if gv_tot is None and gv1 is not None:
            gv_tot = (gv1 or 0) + (gv2 or 0)
        if hv_tot is None and gv_tot is None:
            continue

        year = year_hint
        if isinstance(dt, datetime):
            year = dt.year
        rd = _as_date(dt, year, month)

        def add_r(meter, val, cons):
            nonlocal created, skipped
            if val is None:
                return
            if _reading_exists(db, meter.id, rd):
                skipped += 1
                return
            db.add(
                MeterReading(
                    meter_id=meter.id,
                    reading_date=rd,
                    value_t1=float(val),
                    value_t2=None,
                    consumption_t1=float(cons) if cons is not None else None,
                    consumption_t2=None,
                    source="import",
                )
            )
            created += 1

        add_r(cold, hv_tot, hv_cons)
        add_r(hot, gv_tot, gv_cons)
        add_r(cold1, hv1, None)
        add_r(cold2, hv2, None)
        add_r(hot1, gv1, None)
        add_r(hot2, gv2, None)

    return {"readings_created": created, "readings_skipped": skipped, "meters": ["ХВ", "ГВ"]}


def _import_electricity_sheet(db: Session, ws, property_id: int, year_hint: int) -> dict:
    meter = _get_or_create_meter(
        db, property_id, "electricity", "Электричество Т1/Т2", "day_night", "кВт·ч"
    )
    created = skipped = 0
    for row in ws.iter_rows(min_row=1, values_only=True):
        cells = list(row) if row else []
        if len(cells) < 5:
            continue
        month = None
        col0 = None
        for i, c in enumerate(cells[:3]):
            month = _month_num(c)
            if month:
                col0 = i
                break
        if not month:
            continue
        base = col0 + 1
        dt = cells[base] if base < len(cells) else None
        t1 = _num(cells[base + 1] if base + 1 < len(cells) else None)
        c1 = _num(cells[base + 2] if base + 2 < len(cells) else None)
        # often empty col then night
        t2 = _num(cells[base + 4] if base + 4 < len(cells) else None)
        c2 = _num(cells[base + 5] if base + 5 < len(cells) else None)
        if t2 is None:
            t2 = _num(cells[base + 3] if base + 3 < len(cells) else None)
            c2 = _num(cells[base + 4] if base + 4 < len(cells) else None)
        if t1 is None:
            continue
        year = year_hint
        if isinstance(dt, datetime):
            year = dt.year
        rd = _as_date(dt, year, month)
        if _reading_exists(db, meter.id, rd):
            skipped += 1
            continue
        db.add(
            MeterReading(
                meter_id=meter.id,
                reading_date=rd,
                value_t1=float(t1),
                value_t2=float(t2) if t2 is not None else None,
                consumption_t1=float(c1) if c1 is not None else None,
                consumption_t2=float(c2) if c2 is not None else None,
                source="import",
            )
        )
        created += 1
    return {"readings_created": created, "readings_skipped": skipped, "meters": ["Электричество"]}


def _import_payment_sheet(
    db: Session,
    ws,
    property_id: int,
    year_hint: int,
    provider: str,
    accrued_col: int = 2,
    paid_col: int = 4,
    paid_date_col: int = 5,
) -> dict:
    """Generic payment rows: month, accrued, ..., paid amount, paid date."""
    created = skipped = 0
    for row in ws.iter_rows(min_row=1, values_only=True):
        cells = list(row) if row else []
        if len(cells) < 4:
            continue
        month = None
        col0 = None
        for i, c in enumerate(cells[:3]):
            month = _month_num(c)
            if month:
                col0 = i
                break
        if not month:
            continue
        accrued = _num(cells[accrued_col] if accrued_col < len(cells) else None)
        paid = _num(cells[paid_col] if paid_col < len(cells) else None)
        pdt = cells[paid_date_col] if paid_date_col < len(cells) else None
        if accrued is None and paid is None:
            continue
        if _bill_exists(db, property_id, year_hint, month, provider):
            skipped += 1
            continue
        total_a = float(accrued or 0)
        total_p = float(paid or 0)
        status = "paid" if total_p >= total_a and total_a > 0 else ("partial" if total_p > 0 else "draft")
        bill = Bill(
            property_id=property_id,
            provider=provider,
            period_year=year_hint,
            period_month=month,
            total_accrued=total_a,
            total_paid=total_p,
            status=status,
            notes="Импорт из Excel",
        )
        db.add(bill)
        db.flush()
        if total_p > 0:
            pd = _as_date(pdt, year_hint, month)
            db.add(
                Payment(
                    bill_id=bill.id,
                    amount=total_p,
                    payment_date=pd,
                    method="import",
                )
            )
        created += 1
    return {"bills_created": created, "bills_skipped": skipped, "provider": provider}


def import_workbook(db: Session, property_id: int, file_bytes: bytes) -> dict:
    prop = db.get(Property, property_id)
    if not prop:
        raise ValueError("Объект не найден")

    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    report: dict[str, Any] = {"sheets": [], "summary": {}}
    total_readings = total_bills = 0

    for name in wb.sheetnames:
        ws = wb[name]
        low = name.lower()
        year_hint = 2025
        for token in name.replace("-", " ").split():
            if token.isdigit() and len(token) == 4:
                year_hint = int(token)

        try:
            if "эл" in low or "энерг" in low:
                r = _import_electricity_sheet(db, ws, property_id, year_hint)
                total_readings += r["readings_created"]
                report["sheets"].append({"name": name, "type": "electricity", **r})
            elif low.strip().isdigit() or (low[:4].isdigit() and "оплат" not in low and "еИрц" not in low and "еирц" not in low):
                # year sheet with water
                r = _import_water_sheet(db, ws, property_id, year_hint)
                total_readings += r["readings_created"]
                report["sheets"].append({"name": name, "type": "water", **r})
            elif "техком" in low or (low.startswith("оплат") and "тепло" not in low and "кап" not in low):
                r = _import_payment_sheet(
                    db, ws, property_id, year_hint, "Техкомсервис / ЖКУ",
                    accrued_col=2, paid_col=4, paid_date_col=5,
                )
                total_bills += r["bills_created"]
                report["sheets"].append({"name": name, "type": "payments", **r})
            elif "тепло" in low or "водоканал" in low:
                r = _import_payment_sheet(
                    db, ws, property_id, year_hint, "Теплосеть / Водоканал",
                    accrued_col=7, paid_col=9, paid_date_col=10,
                )
                total_bills += r["bills_created"]
                report["sheets"].append({"name": name, "type": "payments", **r})
            elif "кап" in low or "мусор" in low:
                r = _import_payment_sheet(
                    db, ws, property_id, year_hint, "Капремонт / Мусор",
                    accrued_col=2, paid_col=4, paid_date_col=5,
                )
                total_bills += r["bills_created"]
                report["sheets"].append({"name": name, "type": "payments", **r})
            elif "еирц" in low or "епд" in low:
                report["sheets"].append({
                    "name": name,
                    "type": "eirc_matrix",
                    "note": "Матрица ЕИРЦ: строки услуг по месяцам — используйте OCR квитанций для детализации",
                })
            else:
                report["sheets"].append({"name": name, "type": "skipped", "note": "Неизвестный формат листа"})
        except Exception as e:
            report["sheets"].append({"name": name, "type": "error", "error": str(e)})

    db.flush()
    report["summary"] = {
        "readings_created": total_readings,
        "bills_created": total_bills,
        "property_id": property_id,
        "property_name": prop.name,
    }
    return report

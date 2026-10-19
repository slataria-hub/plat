"""Excel export of user data (readings, bills, payments)."""
from io import BytesIO
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from sqlalchemy import select, desc
from app.models.property import Property
from app.models.meter import Meter, MeterReading
from app.models.bill import Bill, BillItem, Payment, ServiceType

HEADER_FILL = PatternFill("solid", fgColor="0D9488")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Border(
    left=Side(style="thin", color="E5E7EB"),
    right=Side(style="thin", color="E5E7EB"),
    top=Side(style="thin", color="E5E7EB"),
    bottom=Side(style="thin", color="E5E7EB"),
)

METER_LABELS = {
    "cold_water": "Холодная вода",
    "hot_water": "Горячая вода",
    "electricity": "Электричество",
}
MONTHS_RU = [
    "", "январь", "февраль", "март", "апрель", "май", "июнь",
    "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь",
]


def _style_header(ws, row: int, cols: int):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN


def _auto_width(ws):
    for col in ws.columns:
        max_len = 0
        letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[letter].width = min(max_len + 3, 40)


def export_property_to_excel(db: Session, property_id: int, user_id: int) -> BytesIO:
    prop = db.execute(
        select(Property).where(Property.id == property_id, Property.user_id == user_id)
    ).scalar_one_or_none()
    if not prop:
        raise ValueError("Объект не найден")

    wb = Workbook()

    # ── Sheet: Показания ─────────────────────────────────
    ws = wb.active
    ws.title = "Показания"
    headers = ["Счётчик", "Тип", "Дата", "Значение T1", "Значение T2", "Расход T1", "Расход T2", "Источник"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    meters = db.execute(
        select(Meter).where(Meter.property_id == property_id).order_by(Meter.type, Meter.name)
    ).scalars().all()

    for meter in meters:
        readings = db.execute(
            select(MeterReading)
            .where(MeterReading.meter_id == meter.id)
            .order_by(MeterReading.reading_date)
        ).scalars().all()
        for r in readings:
            ws.append([
                meter.name,
                METER_LABELS.get(meter.type, meter.type),
                r.reading_date.isoformat() if r.reading_date else "",
                r.value_t1,
                r.value_t2,
                r.consumption_t1,
                r.consumption_t2,
                r.source,
            ])
    _auto_width(ws)

    # ── Sheet: Квитанции ─────────────────────────────────
    ws2 = wb.create_sheet("Квитанции")
    h2 = ["Период", "Поставщик", "Начислено", "Оплачено", "Статус", "Файл"]
    ws2.append(h2)
    _style_header(ws2, 1, len(h2))

    bills = db.execute(
        select(Bill)
        .where(Bill.property_id == property_id)
        .order_by(desc(Bill.period_year), desc(Bill.period_month))
    ).scalars().all()

    for b in bills:
        period = f"{MONTHS_RU[b.period_month]} {b.period_year}"
        ws2.append([
            period,
            b.provider,
            b.total_accrued,
            b.total_paid,
            b.status,
            b.file_type or "",
        ])
    _auto_width(ws2)

    # ── Sheet: Строки квитанций ──────────────────────────
    ws3 = wb.create_sheet("Строки квитанций")
    h3 = ["Период", "Поставщик", "Услуга", "Объём", "Тариф", "Сумма"]
    ws3.append(h3)
    _style_header(ws3, 1, len(h3))

    service_names = {
        s.id: s.name for s in db.execute(select(ServiceType)).scalars().all()
    }

    for b in bills:
        period = f"{MONTHS_RU[b.period_month]} {b.period_year}"
        for item in b.items:
            name = service_names.get(item.service_type_id) if item.service_type_id else item.raw_name
            ws3.append([period, b.provider, name or "—", item.volume, item.tariff, item.amount])
    _auto_width(ws3)

    # ── Sheet: Оплаты ────────────────────────────────────
    ws4 = wb.create_sheet("Оплаты")
    h4 = ["Период квитанции", "Дата оплаты", "Сумма", "Способ", "Комментарий"]
    ws4.append(h4)
    _style_header(ws4, 1, len(h4))

    for b in bills:
        period = f"{MONTHS_RU[b.period_month]} {b.period_year}"
        for p in b.payments:
            ws4.append([
                period,
                p.payment_date.isoformat() if p.payment_date else "",
                p.amount,
                p.method or "",
                p.notes or "",
            ])
    _auto_width(ws4)

    # ── Sheet: Инфо ──────────────────────────────────────
    ws5 = wb.create_sheet("Объект")
    ws5.append(["Поле", "Значение"])
    _style_header(ws5, 1, 2)
    ws5.append(["Название", prop.name])
    ws5.append(["Адрес", prop.address or ""])
    ws5.append(["Лицевой счёт", prop.account_number or ""])
    ws5.append(["Дата выгрузки", date.today().isoformat()])
    _auto_width(ws5)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
